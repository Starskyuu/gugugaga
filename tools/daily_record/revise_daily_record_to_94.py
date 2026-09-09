from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"F:\gugugaga\06 每日记录\未来工程创新设计 个人过程记录（补充至9.7）.docx")
OUTPUT = Path(r"F:\gugugaga\06 每日记录\未来工程创新设计 个人过程记录（补充至9.4）.docx")

TEXT_94 = (
    "今天继续完善Unity3D三维搜救仿真。为了让演示形成完整的任务闭环，我在场景中布置了60cm×60cm的水域、"
    "22cm×15cm的中心基站、四艘搜救船和可以手动或随机添加的落水者，并把Python规划程序与Unity通过UDP连接。"
    "Unity将船只和人员坐标发送给规划端，接收航点后控制小船沿路线运动。调试时，我重点检查坐标转换、船只初始泊位、"
    "路线方向和到达判定，解决了小船收到路线后不运动的问题。现在小船能够从基站内部出航，接近落水者后完成救援，再返回原泊位；"
    "当外部规划服务没有响应时，系统会切换到本地调度，避免仿真一直停在等待状态。我还调整了船体、人物、基站和水面的视觉效果，"
    "并在界面中显示待救人数、船上人数、已救人数、规划方式和各船任务状态。今天的工作使仿真从静态场景变成了可以连续操作和观察的救援流程。"
)


def set_paragraph_text(paragraph, text):
    style_source = next((r for r in paragraph.runs if r.text), paragraph.runs[0] if paragraph.runs else None)
    rpr = deepcopy(style_source._element.rPr) if style_source is not None and style_source._element.rPr is not None else None
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    run = paragraph.add_run(text)
    if rpr is not None:
        run._element.insert(0, rpr)


doc = Document(SOURCE)
paragraphs = doc.paragraphs

date_94_index = next(i for i, p in enumerate(paragraphs) if p.text.strip() == "9.4")
if date_94_index + 1 >= len(paragraphs):
    raise RuntimeError("9.4正文不存在")
set_paragraph_text(paragraphs[date_94_index + 1], TEXT_94)

paragraphs = doc.paragraphs
for paragraph in paragraphs[date_94_index + 2:]:
    parent = paragraph._element.getparent()
    parent.remove(paragraph._element)

doc.save(OUTPUT)

check = Document(OUTPUT)
remaining_dates = [p.text.strip() for p in check.paragraphs if p.text.strip() in {"9.4", "9.5", "9.6", "9.7"}]
print(f"SAVED={OUTPUT}")
print(f"PARAGRAPHS={len(check.paragraphs)}")
print(f"DATES={remaining_dates}")
