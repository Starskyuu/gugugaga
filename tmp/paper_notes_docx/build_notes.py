from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

ROOT = Path(r"F:\gugugaga")
NOTES = ROOT / "02 reading notes"
OUT = NOTES / "Word版"
OUT.mkdir(parents=True, exist_ok=True)

FILES = [
    "01-自旋转桨轮两栖机器人.md",
    "02-安防机器人自动对接与充电.md",
    "03-模块化空地车队集成设计.md",
    "04-JAXA月球车部署坡道.md",
    "05-月船3号月球车坡道.md",
    "06-视觉-激光雷达自主对接.md",
]

# compact_reference_guide preset, with a named East Asian font override.
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1A1F24"
GRAY = "606770"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
RULE = "C8D0D9"
ASCII_FONT = "Calibri"
EA_FONT = "Microsoft YaHei"

def set_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = ASCII_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ASCII_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), EA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = tcMar.find(qn(f"w:{m}"))
        if el is None:
            el = OxmlElement(f"w:{m}")
            tcMar.append(el)
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_table_widths(table, widths_dxa):
    table.autofit = False
    total = sum(widths_dxa)
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd"); tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120"); tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tcW = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW"); cell._tc.get_or_add_tcPr().append(tcW)
            tcW.set(qn("w:w"), str(widths_dxa[idx])); tcW.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_page_field(paragraph):
    run = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = ASCII_FONT; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1",16,BLUE,18,10), ("Heading 2",13,BLUE,14,7), ("Heading 3",12,DARK_BLUE,10,5)
    ]:
        st = styles[name]
        st.font.name = ASCII_FONT; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
        st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        st = styles[list_name]
        st.font.name = ASCII_FONT; st.font.size = Pt(11)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
        st.paragraph_format.left_indent = Inches(0.375)
        st.paragraph_format.first_line_indent = Inches(-0.188)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.25

    if "Figure Caption" not in styles:
        cap = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        cap = styles["Figure Caption"]
    cap.font.name = ASCII_FONT; cap.font.size = Pt(9); cap.font.italic = True; cap.font.color.rgb = RGBColor.from_string(GRAY)
    cap._element.rPr.rFonts.set(qn("w:eastAsia"), EA_FONT)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(3); cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.keep_with_next = False

def setup_doc(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = Inches(1); sec.bottom_margin = Inches(1); sec.left_margin = Inches(1); sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492); sec.footer_distance = Inches(0.492)
    configure_styles(doc)

    hp = sec.header.paragraphs[0]
    hp.text = "文献阅读笔记  |  水陆两栖应急配送舰队"
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    for r in hp.runs: set_font(r, 8.5, bold=True, color=GRAY)

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("内部学习资料  |  第 "); set_font(r, 8.5, color=GRAY)
    add_page_field(fp)
    r = fp.add_run(" 页"); set_font(r, 8.5, color=GRAY)
    return doc

def clean_inline(s):
    s = re.sub(r"\[打开原文\]\([^)]*\)", "", s)
    s = s.replace("**", "").replace("`", "")
    return s.strip()

def add_rich_text(paragraph, text):
    # Minimal inline Markdown support for bold and code-like emphasis.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part: continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); set_font(r, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); set_font(r, color=DARK_BLUE)
        else:
            r = paragraph.add_run(part); set_font(r)

def add_callout(doc, label, body):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), LIGHT_BLUE); pPr.append(shd)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "160"); ind.set(qn("w:right"), "160"); pPr.append(ind)
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(9); p.paragraph_format.line_spacing = 1.2
    r = p.add_run(label + "  "); set_font(r, 10, bold=True, color=BLUE)
    r = p.add_run(body); set_font(r, 11.5, bold=True, color=INK)

def add_markdown_table(doc, rows):
    if len(rows) < 2: return
    data = []
    for line in rows:
        vals = [clean_inline(v.strip()) for v in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{3,}:?", v.replace(" ", "")) for v in vals):
            continue
        data.append(vals)
    if not data: return
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    if cols == 2: widths = [2700, 6660]
    elif cols == 3: widths = [2100, 3000, 4260]
    else: widths = [9360//cols]*cols; widths[-1] += 9360-sum(widths)
    set_table_widths(table, widths)
    for i,row in enumerate(data):
        for j in range(cols):
            cell = table.cell(i,j)
            cell.text = row[j] if j < len(row) else ""
            if i == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
                for r in p.runs: set_font(r, 9.5, bold=(i==0), color=INK)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)

def add_figure(doc, image_path, alt_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(image_path), width=Inches(4.85))
    docPr = shape._inline.docPr
    docPr.set("descr", alt_text)

def convert(md_path):
    raw = md_path.read_text(encoding="utf-8")
    lines = raw.splitlines()
    title = clean_inline(lines[0].lstrip("# "))
    doc = setup_doc(title)

    # Memo masthead title block.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(5)
    r = p.add_run("文献阅读笔记"); set_font(r, 10, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(8); p.paragraph_format.keep_with_next = True
    r = p.add_run(title); set_font(r, 23, bold=True, color=INK)

    i = 1
    meta_lines = []
    while i < len(lines) and (not lines[i].strip() or lines[i].startswith(">")):
        if lines[i].startswith(">"):
            content = clean_inline(lines[i].lstrip("> "))
            if content: meta_lines.append(content)
        i += 1
    if meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12); p.paragraph_format.line_spacing = 1.15
        r = p.add_run("文献信息  "); set_font(r, 9.5, bold=True, color=BLUE)
        r = p.add_run(" ".join(meta_lines)); set_font(r, 9.5, color=GRAY)

    pending_caption = None
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1; continue
        if stripped.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if m:
                alt, rel = m.groups()
                image_path = (md_path.parent / rel).resolve()
                if image_path.exists(): add_figure(doc, image_path, alt)
                pending_caption = alt
            i += 1; continue
        if stripped.startswith("*图：") and stripped.endswith("*"):
            p = doc.add_paragraph(style="Figure Caption")
            p.add_run(stripped.strip("*"))
            for r in p.runs: set_font(r, 9, italic=True, color=GRAY)
            i += 1; continue
        if stripped.startswith("|"):
            tbl=[]
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip()); i += 1
            add_markdown_table(doc,tbl); continue
        if stripped.startswith("### "):
            doc.add_paragraph(clean_inline(stripped[4:]), style="Heading 2"); i+=1; continue
        if stripped.startswith("## "):
            heading = clean_inline(stripped[3:])
            # Convert the one-sentence section into a lead callout.
            if heading == "一句话概括":
                j=i+1
                while j < len(lines) and not lines[j].strip(): j+=1
                body = clean_inline(lines[j]) if j < len(lines) else ""
                add_callout(doc,"一句话概括",body); i=j+1; continue
            doc.add_paragraph(heading, style="Heading 1"); i+=1; continue
        if re.match(r"^\d+\.\s+", stripped):
            body = re.sub(r"^\d+\.\s+", "", stripped)
            p = doc.add_paragraph(style="List Number"); add_rich_text(p,body); i+=1; continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_rich_text(p,stripped[2:]); i+=1; continue
        if stripped.startswith(">"):
            body=clean_inline(stripped.lstrip("> "))
            if body: add_callout(doc,"原文信息",body)
            i+=1; continue
        # Collect consecutive prose lines into one paragraph.
        para=[stripped]; i+=1
        while i < len(lines):
            nxt=lines[i].strip()
            if not nxt or nxt.startswith(("#","!","|",">","- ")) or re.match(r"^\d+\.\s+",nxt): break
            para.append(nxt); i+=1
        p=doc.add_paragraph(); add_rich_text(p," ".join(para))

    doc.add_paragraph("来源与使用说明", style="Heading 1")
    p=doc.add_paragraph()
    add_rich_text(p,"本笔记依据本地保存的原论文整理；图片均为原文关键页面的本地渲染，仅用于团队内部学习、讨论和设计参考。笔记中的评价与启发属于阅读分析，不等同于作者原文结论。")
    out = OUT / (md_path.stem + ".docx")
    doc.save(out)
    return out

if __name__ == "__main__":
    outputs=[]
    for name in FILES:
        outputs.append(convert(NOTES/name))
    for p in outputs:
        print(p)

