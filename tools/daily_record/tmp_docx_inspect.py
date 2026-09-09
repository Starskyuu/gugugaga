from pathlib import Path
from docx import Document

path = Path(r"F:\gugugaga\06 每日记录\未来工程创新设计 个人过程记录（新版）.docx")
doc = Document(path)
print(f"PATH={path}")
print(f"PARAGRAPHS={len(doc.paragraphs)} TABLES={len(doc.tables)} SECTIONS={len(doc.sections)}")
print("\nPARAGRAPHS")
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\t", "<TAB>").replace("\n", "<BR>")
    if text.strip():
        print(f"P{i:04d} [{p.style.name}] {text}")
print("\nTABLES")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else ''}")
    for ri, row in enumerate(table.rows):
        cells = [c.text.replace("\n", "<BR>") for c in row.cells]
        print(f"T{ti}R{ri}: {cells}")
