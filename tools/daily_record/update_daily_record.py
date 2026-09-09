from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph


SOURCE = Path(r"F:\gugugaga\06 每日记录\未来工程创新设计 个人过程记录（新版）.docx")
OUTPUT = Path(r"F:\gugugaga\06 每日记录\未来工程创新设计 个人过程记录（补充至9.7）.docx")

ENTRIES = {
    "9.3": (
        "今天把前一天的软件设想落实到Unity场景中。我先按课程模型的尺度搭建救援水域、中心基站和四艘两栖搜救船，"
        "再把落水者位置、船只位置和规划路线放进同一界面。为了让演示不只停留在静态模型，我继续整理坐标转换、任务分配"
        "和航点跟随逻辑，使小船能够按照规划结果执行任务。调试过程中，我逐项检查场景比例、相机视角、船只初始位置和路线显示，"
        "并尝试把A*路径规划与Unity中的运动控制衔接起来。这个过程让我意识到，算法给出一串坐标并不等于仿真已经完成，坐标系方向、"
        "模型尺寸和运动状态都可能让小船偏离预期。今天完成了第一版可交互仿真框架，也为后续加入救援动作、返航和多船协同打下了基础。"
    ),
    "9.4": (
        "今天主要补充项目的理论依据和现实工程分析。我阅读并整理了六篇与两栖移动、自动对接、模块化救援装备、折叠坡道和视觉定位有关的论文，"
        "分别提炼研究问题、机构方案、实验结果和局限，并制作成带原文配图的阅读笔记。结合这些文献，我进一步梳理了城市方舟式救援基站在现实建设中"
        "可能遇到的难题，包括防洪与结构安全、供电与储能、通信中继、船只回收对接、设备维护和多系统联调。我还单独复习了A*算法，弄清启发式函数"
        "如何影响搜索效率，以及在实际救援场景中还要加入障碍膨胀、安全距离和动态重规划。今天的工作让我发现，课程模型可以用简化机构展示核心思路，"
        "但现实方案必须说明环境条件、可靠性和维护方式，不能把论文中的样机数据直接当作本项目指标。"
    ),
    "9.5": (
        "今天把前一天的六篇阅读笔记汇总成一份综合稿，并根据项目对象重新调整表述。由于我们的核心设备是水陆两栖船，我将论文中的机器人和月球车案例"
        "改写为可供两栖船与基站设计借鉴的经验，重点说明自动返航、末端停靠、坡道部署和补能流程，同时删去生硬的概念堆叠，使文稿更像自己的研究记录。"
        "下午继续完善Unity仿真，重点解决小船启动后不运动的问题，并补上救援闭环。现在小船能够从基站内部泊位出发，接近落水者后将其带上船，再返回原泊位；"
        "当外部Python规划服务没有响应时，系统也能切换到本地调度。与此同时，我调整了船体、人物、基站和控制面板的视觉效果，让任务状态、待救人数、"
        "船上人数和已救人数能够直接显示。"
    ),
    "9.6": (
        "今天继续围绕多船救援的可用性改进程序。我在原有单人快返模式之外增加了多人连续救援模式，每艘船可以依次接近多名落水者，达到四人容量或现场无人待救后再返航，"
        "界面可以在两种模式之间切换。为了观察不同任务规模下的表现，我把随机生成落水者和重复派单逻辑接入同一套管理程序，并将人数上限扩展到一千。"
        "多船同时运动后出现了路线交叉和船体靠得过近的问题，因此我加入实时避碰转向、最小安全距离和碰撞盒预测，船头相向时会向右避让。每救起一人后，"
        "系统会根据剩余目标重新安排路线。经过Unity重新编译和运行测试，救援人数能够持续累加，目标会逐步清零，运行日志中没有出现新的异常。"
    ),
    "9.7": (
        "今天把仿真从功能原型继续推进到可用于展示和后续开发的模型工程。我使用Blender制作并整理了基站、树莓派、俯视相机、搜救船和落水者的白模，统一命名、比例和导出设置，"
        "再将FBX模型放入Unity资源目录，由程序自动替换原来的基础几何体。针对救援动作，我继续调整人物模型的尺寸和登船位置，使人物被救起后能够随船移动，并为多人模式预留不同座位。"
        "我还检查了基站内部泊位、船只出入口和相机覆盖范围，确保模型布局与软件中的坐标系统一致。今天的工作让我更清楚地看到，最终仿真应用不仅需要算法正确，"
        "还要建立稳定的模型资源规范，否则后续每次替换模型都可能破坏碰撞范围、登船位置或路线显示。"
    ),
}


def copy_run_style(source_run, target_run):
    if source_run is not None and source_run._element.rPr is not None:
        target_run._element.insert(0, deepcopy(source_run._element.rPr))


def set_paragraph_text(paragraph, text, source_run):
    for run in list(paragraph.runs):
        paragraph._element.remove(run._element)
    run = paragraph.add_run(text)
    copy_run_style(source_run, run)


def add_after(paragraph, text, style_source, run_source):
    new_p = deepcopy(style_source._element)
    for child in list(new_p):
        if child.tag.endswith("}r") or child.tag.endswith("}hyperlink"):
            new_p.remove(child)
    paragraph._element.addnext(new_p)
    wrapper = Paragraph(new_p, paragraph._parent)
    set_paragraph_text(wrapper, text, run_source)
    return wrapper


doc = Document(SOURCE)
paragraphs = doc.paragraphs

date_template = paragraphs[8]
body_template = paragraphs[7]
date_run = next((r for r in date_template.runs if r.text), date_template.runs[0] if date_template.runs else None)
body_run = next((r for r in body_template.runs if r.text), body_template.runs[0] if body_template.runs else None)

# Expand the existing short 9.3 note.
set_paragraph_text(paragraphs[9], ENTRIES["9.3"], body_run)

# The source already ends with the 9.4 date. Fill it, then append later dates.
cursor = doc.paragraphs[-1]
if cursor.text.strip() != "9.4":
    raise RuntimeError(f"Expected final source date 9.4, found: {cursor.text!r}")
cursor = add_after(cursor, ENTRIES["9.4"], body_template, body_run)

for date in ("9.5", "9.6", "9.7"):
    cursor = add_after(cursor, date, date_template, date_run)
    cursor = add_after(cursor, ENTRIES[date], body_template, body_run)

doc.save(OUTPUT)
print(f"SAVED={OUTPUT}")
print(f"PARAGRAPHS={len(doc.paragraphs)}")
